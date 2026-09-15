import json
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock
import openpyxl
import pytest

from backend.agent.engine import AgentEngine
from backend.agent.memory import ConversationMemory
from backend.config import settings
from backend.tools.xlsx_report import create_xlsx_report, XlsxReportInput
from backend.tools.artifact_verifier import create_artifact_verifier, ArtifactVerifierInput


def _make_engine():
    memory = ConversationMemory()
    router = MagicMock()
    return AgentEngine(settings=settings, router=router, memory=memory)


@pytest.mark.asyncio
async def test_synthesize_xlsx_data_from_reasoning():
    engine = _make_engine()
    mock_provider = MagicMock()
    mock_resp = MagicMock()
    mock_resp.content = json.dumps({
        "headers": ["Equipment Tag", "Component", "Condition", "Action Required"],
        "rows": [
            ["P-204", "Mechanical Seal", "Leaking hydrocarbon stream", "Replace O-ring"],
            ["P-204", "Bearings", "Elevated vibration (7.2 mm/s)", "Lubricate and rebalance"],
        ]
    })
    mock_provider.chat = AsyncMock(return_value=mock_resp)

    executed_step_results = [
        {
            "tool": "reasoning",
            "description": "Synthesize inspection observations",
            "result": "P-204 pump seal is leaking and bearing vibration is critical at 7.2 mm/s.",
        }
    ]

    result = await engine._synthesize_xlsx_data(
        user_request="Prepare an Excel report of P-204 anomalies",
        filename="P204_Anomalies.xlsx",
        step_description="Generate audit spreadsheet",
        executed_step_results=executed_step_results,
        sources=[],
        provider=mock_provider,
        model_name="test-model",
    )

    assert "headers" in result
    assert "rows" in result
    assert len(result["headers"]) == 4
    assert len(result["rows"]) == 2
    assert result["rows"][0][0] == "P-204"


@pytest.mark.asyncio
async def test_synthesize_xlsx_data_fallback():
    engine = _make_engine()
    mock_provider = MagicMock()
    # Simulate LLM failure or malformed JSON
    mock_provider.chat = AsyncMock(side_effect=RuntimeError("LLM unavailable"))

    executed_step_results = [
        {
            "tool": "document_search",
            "description": "Search runbook",
            "result": "Found P-204 runbook section 4.2",
        }
    ]

    result = await engine._synthesize_xlsx_data(
        user_request="Export data table",
        filename="export.xlsx",
        step_description="Write spreadsheet",
        executed_step_results=executed_step_results,
        sources=[],
        provider=mock_provider,
        model_name="test-model",
    )

    assert "headers" in result
    assert "rows" in result
    assert len(result["rows"]) >= 1
    assert result["headers"] == ["Item", "Details"]


@pytest.mark.asyncio
async def test_xlsx_report_writes_populated_data(tmp_path: Path):
    execute_xlsx = create_xlsx_report(tmp_path)
    report_input = XlsxReportInput(
        filename="compliance_audit.xlsx",
        title="MRPL Refinery Environmental Compliance",
        headers=["Parameter", "Standard", "Measured", "Compliance"],
        rows=[
            ["SO2 Emission", "< 50 mg/Nm3", "32 mg/Nm3", "PASS"],
            ["NOx Emission", "< 150 mg/Nm3", "120 mg/Nm3", "PASS"],
        ],
        summary_notes="All atmospheric emissions within CPCB norms.",
    )

    res = await execute_xlsx(report_input)
    assert "created_path" in res
    assert "sha256_hash" in res
    assert res["row_count"] == 2

    wb_path = tmp_path / "compliance_audit.xlsx"
    assert wb_path.exists()

    wb = openpyxl.load_workbook(wb_path)
    ws = wb.active
    assert ws is not None

    # Verify headers and data rows exist in the spreadsheet
    found_headers = False
    found_rows = 0
    for row in ws.iter_rows(values_only=True):
        if row and "Parameter" in row and "Standard" in row:
            found_headers = True
        if row and "SO2 Emission" in row:
            found_rows += 1
        if row and "NOx Emission" in row:
            found_rows += 1

    assert found_headers is True
    assert found_rows == 2


@pytest.mark.asyncio
async def test_artifact_verifier_validates_populated_xlsx(tmp_path: Path):
    execute_xlsx = create_xlsx_report(tmp_path)
    await execute_xlsx(XlsxReportInput(
        filename="equipment_status.xlsx",
        title="Equipment Status Report",
        headers=["Tag", "Status", "Pressure"],
        rows=[["P-204", "Active", "42 bar"]],
    ))

    verify_fn = create_artifact_verifier(tmp_path)
    v_res = await verify_fn(ArtifactVerifierInput(
        relative_path="equipment_status.xlsx",
        expected_columns=["Tag", "Status"],
        min_row_count=1,
    ))

    assert v_res["verified"] is True
    assert v_res["row_count"] >= 1
    assert "sha256_hash" in v_res


@pytest.mark.asyncio
async def test_artifact_verifier_relative_path_resolution(tmp_path: Path):
    execute_xlsx = create_xlsx_report(tmp_path)
    await execute_xlsx(XlsxReportInput(
        filename="verified_report.xlsx",
        title="Verified Report",
        headers=["Col1", "Col2"],
        rows=[["Val1", "Val2"]],
    ))

    verify_fn = create_artifact_verifier(tmp_path)
    # Test path resolution with data/sandbox/ prefix
    v_res = await verify_fn(ArtifactVerifierInput(
        relative_path="data/sandbox/verified_report.xlsx",
        min_row_count=1,
    ))
    assert v_res["verified"] is True


@pytest.mark.asyncio
async def test_xlsx_end_to_end_pipeline(tmp_path: Path):
    engine = _make_engine()
    execute_xlsx = create_xlsx_report(tmp_path)
    verify_fn = create_artifact_verifier(tmp_path)

    # 1. Mock synthesis
    mock_provider = MagicMock()
    mock_provider.chat = AsyncMock(return_value=MagicMock(content=json.dumps({
        "headers": ["Tag", "Parameter", "Value"],
        "rows": [["P-204", "Discharge Pressure", "42 bar"]]
    })))

    synthesized = await engine._synthesize_xlsx_data(
        user_request="Create P-204 report",
        filename="P204_Pipeline.xlsx",
        step_description="Generate report",
        executed_step_results=[{"tool": "document_search", "result": "Discharge pressure 42 bar"}],
        sources=[],
        provider=mock_provider,
        model_name="test-model",
    )

    # 2. Write spreadsheet using synthesized data
    xlsx_res = await execute_xlsx(XlsxReportInput(
        filename="P204_Pipeline.xlsx",
        title="P-204 Hydrocracker Report",
        headers=synthesized["headers"],
        rows=synthesized["rows"],
    ))
    assert "created_path" in xlsx_res

    # 3. Verify artifact
    v_res = await verify_fn(ArtifactVerifierInput(
        relative_path="P204_Pipeline.xlsx",
        expected_columns=["Tag", "Parameter"],
        min_row_count=1,
    ))
    assert v_res["verified"] is True


def test_xlsx_workbook_content_integrity(tmp_path: Path):
    import asyncio
    execute_xlsx = create_xlsx_report(tmp_path)
    asyncio.run(execute_xlsx(XlsxReportInput(
        filename="integrity_test.xlsx",
        title="Integrity Inspection Report",
        headers=["Component", "Spec", "Observed"],
        rows=[
            ["Impeller", "316 SS", "Pass"],
            ["Casing", "Carbon Steel", "Pass"],
        ],
    )))

    wb = openpyxl.load_workbook(tmp_path / "integrity_test.xlsx")
    ws = wb.active
    assert ws is not None

    # Title is in row 1
    assert ws.cell(row=1, column=1).value == "Integrity Inspection Report"
    # Row 4 has headers
    assert ws.cell(row=4, column=1).value == "Component"
    assert ws.cell(row=4, column=2).value == "Spec"
    assert ws.cell(row=4, column=3).value == "Observed"
    # Row 5 has first data row
    assert ws.cell(row=5, column=1).value == "Impeller"
    assert ws.cell(row=5, column=2).value == "316 SS"
    assert ws.cell(row=5, column=3).value == "Pass"


@pytest.mark.asyncio
async def test_synthesize_xlsx_data_extracts_real_p204_evidence():
    """
    Test 1:
    _synthesize_xlsx_data() extracts actual values from retrieved P-204 evidence
    and does NOT emit 'Not stated in retrieved document.' for fields that exist.
    """
    engine = _make_engine()
    doc_path = Path("data/demo/pump_p204_maintenance.md")
    assert doc_path.exists()
    doc_text = doc_path.read_text(encoding="utf-8")

    executed_step_results = [
        {
            "tool": "document_search",
            "description": "Search P-204 maintenance findings and recommendations",
            "result": [
                {
                    "filename": "pump_p204_maintenance.md",
                    "text": doc_text,
                    "page": 1,
                }
            ],
        }
    ]

    # Mock provider returning the requested columns with placeholder or 'Not stated'
    # which synthesis must verify and ground using context evidence
    mock_provider = MagicMock()
    mock_provider.chat = AsyncMock(return_value=MagicMock(content=json.dumps({
        "headers": ["Equipment ID", "Maintenance Findings", "Operating Observations", "Recommended Actions"],
        "rows": [
            [
                "P-204 (Boiler Feed Water Multi-stage Centrifugal Pump Train B)",
                "Not stated in retrieved document.",
                "Not stated in retrieved document.",
                "Not stated in retrieved document.",
            ]
        ]
    })))

    result = await engine._synthesize_xlsx_data(
        user_request=(
            "Create a real P-204 XLSX maintenance data report using the indexed documents. "
            "Include the relevant equipment ID, maintenance findings, operating observations, and recommended actions."
        ),
        filename="P204_Maintenance_Report.xlsx",
        step_description="Synthesize structured maintenance data spreadsheet",
        executed_step_results=executed_step_results,
        sources=[],
        provider=mock_provider,
        model_name="test-model",
    )

    assert "headers" in result
    assert "rows" in result
    headers = result["headers"]
    row = result["rows"][0]

    assert "Equipment ID" in headers
    assert "Maintenance Findings" in headers
    assert "Operating Observations" in headers
    assert "Recommended Actions" in headers

    # Equipment ID must be grounded
    assert "P-204" in str(row[headers.index("Equipment ID")])

    # Maintenance Findings must contain real evidence (not 'Not stated')
    findings_val = str(row[headers.index("Maintenance Findings")])
    assert "not stated in retrieved document" not in findings_val.lower()
    assert any(term in findings_val.lower() for term in ("bearing", "alarm", "temperature", "cavitation", "strainer", "impeller"))

    # Operating Observations must contain real evidence (not 'Not stated')
    obs_val = str(row[headers.index("Operating Observations")])
    assert "not stated in retrieved document" not in obs_val.lower()
    assert any(term in obs_val.lower() for term in ("pressure", "temperature", "cavitation", "bar", "88.4"))

    # Recommended Actions must contain real evidence (not 'Not stated')
    actions_val = str(row[headers.index("Recommended Actions")])
    assert "not stated in retrieved document" not in actions_val.lower()
    assert any(term in actions_val.lower() for term in ("impeller", "bearing", "seal", "strainer", "logging", "monitoring"))


@pytest.mark.asyncio
async def test_not_stated_only_used_when_field_truly_absent():
    """
    Test 2 & 3:
    'Not stated in retrieved document.' is ONLY used when the requested field
    truly has no evidence in the document (e.g. OEM Warranty Expiration Date).
    """
    engine = _make_engine()
    doc_path = Path("data/demo/pump_p204_maintenance.md")
    doc_text = doc_path.read_text(encoding="utf-8")

    executed_step_results = [
        {
            "tool": "document_search",
            "description": "Search P-204 data",
            "result": [{"filename": "pump_p204_maintenance.md", "text": doc_text}],
        }
    ]

    mock_provider = MagicMock()
    mock_provider.chat = AsyncMock(return_value=MagicMock(content=json.dumps({
        "headers": ["Equipment ID", "Maintenance Findings", "OEM Warranty Expiration Date"],
        "rows": [
            [
                "P-204 (Boiler Feed Water Multi-stage Centrifugal Pump Train B)",
                "Not stated in retrieved document.",
                "Not stated in retrieved document.",
            ]
        ]
    })))

    result = await engine._synthesize_xlsx_data(
        user_request="Include equipment ID, maintenance findings, and OEM warranty expiration date.",
        filename="P204_Warranty.xlsx",
        step_description="Generate spreadsheet",
        executed_step_results=executed_step_results,
        sources=[],
        provider=mock_provider,
        model_name="test-model",
    )

    headers = result["headers"]
    row = result["rows"][0]

    # Maintenance Findings must have real evidence
    findings_val = str(row[headers.index("Maintenance Findings")])
    assert "not stated in retrieved document" not in findings_val.lower()

    # OEM Warranty Expiration Date MUST say 'Not stated in retrieved document.'
    warranty_val = str(row[headers.index("OEM Warranty Expiration Date")])
    assert "not stated in retrieved document" in warranty_val.lower()


@pytest.mark.asyncio
async def test_generated_xlsx_contains_meaningful_p204_maintenance_data(tmp_path: Path):
    """
    Test 4:
    The generated workbook contains meaningful P-204 maintenance information
    written to the actual spreadsheet cells.
    """
    execute_xlsx = create_xlsx_report(tmp_path)
    engine = _make_engine()
    doc_text = Path("data/demo/pump_p204_maintenance.md").read_text(encoding="utf-8")

    synthesized = await engine._synthesize_xlsx_data(
        user_request="Create P-204 maintenance report with equipment ID, findings, observations, and actions.",
        filename="P204_Real_Report.xlsx",
        step_description="Generate report",
        executed_step_results=[{"tool": "document_search", "result": [{"filename": "pump_p204_maintenance.md", "text": doc_text}]}],
        sources=[],
        provider=MagicMock(chat=AsyncMock(return_value=MagicMock(content="invalid json"))),  # tests evidence fallback
        model_name="test-model",
    )

    res = await execute_xlsx(XlsxReportInput(
        filename="P204_Real_Report.xlsx",
        title="P-204 Boiler Feed Water Pump Maintenance Report",
        headers=synthesized["headers"],
        rows=synthesized["rows"],
    ))

    wb_path = tmp_path / "P204_Real_Report.xlsx"
    assert wb_path.exists()

    wb = openpyxl.load_workbook(wb_path)
    ws = wb.active

    # Check cell values
    all_text = " ".join(str(cell.value) for row in ws.iter_rows() for cell in row if cell.value is not None)
    assert "P-204" in all_text
    assert any(term in all_text.lower() for term in ("bearing", "alarm", "cavitation", "pressure", "impeller"))
    assert "not stated in retrieved document" not in all_text.lower()


@pytest.mark.asyncio
async def test_artifact_verifier_detects_actual_workbook_structure_and_content(tmp_path: Path):
    """
    Test 5:
    artifact_verifier reads the actual generated workbook and confirms:
    - correct filename
    - worksheet exists
    - headers exist
    - at least one populated data row exists
    - expected P-204 values are present
    """
    execute_xlsx = create_xlsx_report(tmp_path)
    verify_fn = create_artifact_verifier(tmp_path)

    # Generate workbook with P-204 evidence
    await execute_xlsx(XlsxReportInput(
        filename="P204_Verified.xlsx",
        title="P-204 Pump Maintenance Diligence",
        headers=["Equipment ID", "Maintenance Findings", "Operating Observations", "Recommended Actions"],
        rows=[
            [
                "P-204 (Boiler Feed Water Multi-stage Centrifugal Pump Train B)",
                "DE radial bearing high-temperature alarm (peak 88.4°C); Suction strainer magnetite clogging.",
                "Audible high-frequency cavitation noise; Discharge pressure dropped to 54 bar.",
                "Installed new 13Cr impeller; Implemented daily delta-P logging.",
            ]
        ],
    ))

    # Verify workbook
    v_res = await verify_fn(ArtifactVerifierInput(
        relative_path="P204_Verified.xlsx",
        expected_columns=["Equipment ID", "Maintenance Findings", "Operating Observations", "Recommended Actions"],
        expected_content=["P-204", "88.4°C", "cavitation", "impeller"],
        min_row_count=1,
    ))

    assert v_res["verified"] is True
    assert v_res["row_count"] == 1
    assert len(v_res["detected_headers"]) == 4
    assert "P204_Verified.xlsx" in v_res["filename"]


@pytest.mark.asyncio
async def test_e2e_document_search_synthesis_xlsx_verifier_pipeline(tmp_path: Path):
    """
    Test 6:
    End-to-end pipeline:
    document_search → synthesis → xlsx_report → artifact_verifier
    produces a populated verified workbook.
    """
    engine = _make_engine()
    execute_xlsx = create_xlsx_report(tmp_path)
    verify_fn = create_artifact_verifier(tmp_path)

    doc_text = Path("data/demo/pump_p204_maintenance.md").read_text(encoding="utf-8")

    # Step 1: document_search
    executed_steps = [
        {
            "tool": "document_search",
            "description": "Retrieve P-204 maintenance findings",
            "result": [{"filename": "pump_p204_maintenance.md", "text": doc_text}],
        }
    ]

    # Step 2: synthesis
    synthesized = await engine._synthesize_xlsx_data(
        user_request="Create P-204 maintenance spreadsheet with equipment ID, findings, observations, and actions.",
        filename="P204_E2E_Report.xlsx",
        step_description="Generate Excel report",
        executed_step_results=executed_steps,
        sources=[],
        provider=MagicMock(chat=AsyncMock(return_value=MagicMock(content=json.dumps({
            "table": [
                ["Equipment ID", "Maintenance Findings", "Operating Observations", "Recommended Actions"],
                ["P-204", "Not stated in retrieved document.", "Cavitation noise", "Replace impeller"],
            ]
        })))),
        model_name="test-model",
    )

    # Step 3: xlsx_report
    report_res = await execute_xlsx(XlsxReportInput(
        filename="P204_E2E_Report.xlsx",
        title="P-204 Maintenance Inspection Report",
        headers=synthesized["headers"],
        rows=synthesized["rows"],
    ))
    assert report_res["row_count"] == 1

    # Step 4: artifact_verifier
    verifier_res = await verify_fn(ArtifactVerifierInput(
        relative_path="P204_E2E_Report.xlsx",
        expected_columns=["Equipment ID", "Maintenance Findings"],
        expected_content=["P-204", "cavitation"],
        min_row_count=1,
    ))
    assert verifier_res["verified"] is True
    assert verifier_res["row_count"] >= 1

