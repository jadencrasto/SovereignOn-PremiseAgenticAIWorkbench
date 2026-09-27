"""
tests/backend/test_xlsx_relationship_workflow.py
-------------------------------------------------
Regression tests for:
1. Relationship header extraction (Problem / Recommended Improvement).
2. Generic evidence-grounded row extraction without hardcoded documents or problems.
3. Filtering out unrelated RAG sources and file_list metadata when file_read provides document content.
4. Schema normalization when LLM produces generic Topic|Description or Page Number|Content.
5. End-to-end spreadsheet creation and verification for problem / recommended improvement requests.
"""

import json
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock
import openpyxl
import pytest

from backend.agent.engine import AgentEngine
from backend.agent.memory import ConversationMemory
from backend.config import settings
from backend.tools.xlsx_report import create_xlsx_report, XlsxReportInput
from backend.tools.artifact_verifier import create_artifact_verifier, ArtifactVerifierInput


def _make_engine():
    memory = ConversationMemory()
    router = MagicMock()
    return AgentEngine(settings=settings, router=router, memory=memory)


# ===========================================================================
# 1. Header Extraction: Relationship Detection
# ===========================================================================

def test_extract_explicit_headers_problem_and_recommended_improvement():
    """When the user asks for problems and recommended improvements, headers are ['Problem', 'Recommended Improvement']."""
    queries = [
        "Create an Excel spreadsheet summarizing the pump problems mentioned in the document and the recommended improvements for each problem.",
        "Summarize the problems mentioned in the document and the recommended improvements for each problem.",
        "Create an Excel report listing pump problems and recommended improvements.",
        "Generate a spreadsheet of the equipment problems and improvements.",
        "Include a table with problems and recommended improvements.",
    ]
    for q in queries:
        headers = AgentEngine._extract_explicit_requested_headers(q)
        assert headers == ["Problem", "Recommended Improvement"], f"Failed for query: {q}"


def test_extract_explicit_headers_other_relationships():
    """Verify other relationship pairs (issues/solutions, risks/mitigations)."""
    assert AgentEngine._extract_explicit_requested_headers("Summarize the issues and solutions") == ["Issue", "Solution"]
    assert AgentEngine._extract_explicit_requested_headers("Create an Excel sheet of risks and mitigations") == ["Risk", "Mitigation"]


def test_extract_explicit_headers_preserves_maintenance_and_generic():
    """Existing maintenance 4-column schema and generic requests are preserved."""
    maint_req = (
        "Create an Excel spreadsheet including the equipment ID, maintenance findings, "
        "operating observations, and recommended actions."
    )
    assert AgentEngine._extract_explicit_requested_headers(maint_req) == [
        "Equipment ID", "Maintenance Findings", "Operating Observations", "Recommended Actions"
    ]

    generic_req = "Create an Excel summary of Q3 chemical lab testing parameters."
    assert AgentEngine._extract_explicit_requested_headers(generic_req) == []


# ===========================================================================
# 2. Relationship Row Extraction from Document Context
# ===========================================================================

def test_extract_relationship_tabular_rows_from_structured_markdown():
    """Deterministically extracts paired problems and improvements from markdown with section headers."""
    doc_path = Path("data/demo/pump_p204_maintenance.md")
    assert doc_path.exists()
    doc_text = doc_path.read_text(encoding="utf-8")

    rows = AgentEngine._extract_relationship_tabular_rows(doc_text, "Problem", "Recommended Improvement")
    assert len(rows) >= 3
    # Check that row 0 has a problem and an improvement
    p0, i0 = rows[0]
    assert any(k in p0.lower() for k in ("bearing", "alarm", "temperature", "strainer", "cavitation", "impeller"))
    assert any(k in i0.lower() for k in ("skf", "bearing", "strainer", "impeller", "install", "clean", "flush", "seal"))


def test_extract_relationship_tabular_rows_from_unstructured_text():
    """Extracts paired problems and improvements from unstructured engineering text without markdown headers."""
    sample_text = (
        "Rotor dynamic behavior of high energy centrifugal pumps is significantly affected by fluid forces. "
        "Fluid force excitation in annular seals generates self-excited sub-synchronous vibration and rotor whirl. "
        "To mitigate this instability, implementing swirl brakes at the seal inlet and optimizing running clearances is required. "
        "Furthermore, internal recirculation at off-design low flow causes severe impeller cavitation and pressure pulsations. "
        "Pumps should be operated strictly above minimum continuous stable flow and suction piping geometry should be modified."
    )
    rows = AgentEngine._extract_relationship_tabular_rows(sample_text, "Problem", "Recommended Improvement")
    assert len(rows) >= 2
    assert any("whirl" in r[0].lower() or "vibration" in r[0].lower() or "fluid force" in r[0].lower() for r in rows)
    assert any("swirl break" in r[1].lower() or "clearance" in r[1].lower() or "flow" in r[1].lower() for r in rows)


# ===========================================================================
# 3. Context Filtering: No Unrelated RAG Sources when file_read Exists
# ===========================================================================

@pytest.mark.asyncio
async def test_synthesize_xlsx_data_uses_file_read_and_filters_unrelated_sources():
    """When file_read content is present, unrelated RAG sources and file_list metadata are filtered out."""
    engine = _make_engine()
    mock_provider = MagicMock()
    mock_resp = MagicMock()
    mock_resp.content = json.dumps({
        "headers": ["Problem", "Recommended Improvement"],
        "rows": [
            ["Sub-synchronous vibration from fluid forces", "Install swirl brakes and optimize seal clearances"],
            ["Impeller cavitation pitting erosion", "Install 13Cr stainless steel impeller and maintain stable flow"],
        ]
    })
    mock_provider.chat = AsyncMock(return_value=mock_resp)

    executed_step_results = [
        {
            "tool": "file_list",
            "description": "List files",
            "result": {"files": [{"filename": "unrelated_doc.txt", "size_bytes": 100}]},
        },
        {
            "tool": "file_read",
            "description": "Read pump document",
            "result": {
                "filename": "Pump Instability Phenomena.pdf",
                "content": "Rotor dynamic fluid forces cause severe sub-synchronous vibration. Install swirl brakes.",
            },
        },
    ]

    # Unrelated RAG sources from a completely different document
    unrelated_source = MagicMock()
    unrelated_source.filename = "unrelated_refinery_benchmark.pdf"
    unrelated_source.page = 1
    unrelated_source.text = "Refinery benchmark SOP 2024: general facility safety protocols."

    res = await engine._synthesize_xlsx_data(
        user_request="Create an Excel spreadsheet summarizing the pump problems mentioned in the document and the recommended improvements for each problem.",
        filename="Pump_Problems_and_Improvements.xlsx",
        step_description="Generate spreadsheet",
        executed_step_results=executed_step_results,
        sources=[unrelated_source],
        provider=mock_provider,
        model_name="test-model",
    )

    assert res["headers"] == ["Problem", "Recommended Improvement"]
    assert len(res["rows"]) == 2
    # Verify the prompt sent to LLM contains the file_read content and NOT the unrelated RAG source
    call_args = mock_provider.chat.call_args[0][0]
    user_prompt_sent = call_args.messages[1].content
    assert "Pump Instability Phenomena.pdf" in user_prompt_sent or "sub-synchronous vibration" in user_prompt_sent
    assert "unrelated_refinery_benchmark.pdf" not in user_prompt_sent


# ===========================================================================
# 4. Schema Normalization: Mapping Topic|Description to Problem|Improvement
# ===========================================================================

@pytest.mark.asyncio
async def test_synthesize_xlsx_data_normalizes_generic_topic_description_to_problem_improvement():
    """If the LLM returns generic headers ['Topic', 'Description'], normalize to ['Problem', 'Recommended Improvement']."""
    engine = _make_engine()
    mock_provider = MagicMock()
    mock_provider.chat = AsyncMock(return_value=MagicMock(content=json.dumps({
        "headers": ["Topic", "Description"],
        "rows": [
            [
                "Sub-synchronous rotor vibration",
                "Install swirl brakes at wear ring inlet and adjust radial clearances",
            ],
            [
                "Impeller cavitation erosion",
                "Operate above minimum continuous flow and upgrade to 13Cr martensitic stainless steel",
            ],
        ]
    })))

    executed_step_results = [
        {
            "tool": "file_read",
            "result": {
                "filename": "Pump_Analysis.txt",
                "content": "Pump vibration and cavitation issues observed. Recommendations provided.",
            },
        }
    ]

    res = await engine._synthesize_xlsx_data(
        user_request="Create an Excel spreadsheet summarizing the pump problems mentioned in the document and the recommended improvements for each problem.",
        filename="Pump_Report.xlsx",
        step_description="Create spreadsheet",
        executed_step_results=executed_step_results,
        sources=[],
        provider=mock_provider,
        model_name="test-model",
    )

    assert res["headers"] == ["Problem", "Recommended Improvement"]
    assert len(res["rows"]) == 2
    assert "vibration" in res["rows"][0][0].lower()
    assert "swirl" in res["rows"][0][1].lower() and "brake" in res["rows"][0][1].lower()


# ===========================================================================
# 5. Fallback Extraction on LLM Failure
# ===========================================================================

@pytest.mark.asyncio
async def test_synthesize_xlsx_data_fallback_for_problems_and_improvements():
    """When the LLM fails or is unavailable, fallback extracts populated Problem and Recommended Improvement rows."""
    engine = _make_engine()
    mock_provider = MagicMock()
    mock_provider.chat = AsyncMock(side_effect=RuntimeError("LLM unavailable"))

    doc_path = Path("data/demo/pump_p204_maintenance.md")
    doc_text = doc_path.read_text(encoding="utf-8")

    executed_step_results = [
        {
            "tool": "file_read",
            "result": {"filename": "pump_p204_maintenance.md", "content": doc_text},
        }
    ]

    res = await engine._synthesize_xlsx_data(
        user_request="Create an Excel spreadsheet summarizing the pump problems mentioned in the document and the recommended improvements for each problem.",
        filename="Pump_Problems.xlsx",
        step_description="Create spreadsheet",
        executed_step_results=executed_step_results,
        sources=[],
        provider=mock_provider,
        model_name="test-model",
    )

    assert res["headers"] == ["Problem", "Recommended Improvement"]
    assert len(res["rows"]) >= 2
    for row in res["rows"]:
        assert len(row) == 2
        assert row[0] != "Not stated in retrieved document."
        assert row[1] != "Not stated in retrieved document."
        assert not row[0].startswith("[")


# ===========================================================================
# 6. End-to-End Workbook Creation & Verification
# ===========================================================================

@pytest.mark.asyncio
async def test_e2e_xlsx_report_and_verifier_for_problem_and_improvement(tmp_path: Path):
    """End-to-end test: xlsx_report generates workbook with Problem | Recommended Improvement and artifact_verifier passes."""
    execute_xlsx = create_xlsx_report(tmp_path)
    verify_artifact = create_artifact_verifier(tmp_path)

    filename = "Pump_Problems_and_Recommended_Improvements.xlsx"
    headers = ["Problem", "Recommended Improvement"]
    rows = [
        [
            "Sub-synchronous lateral rotor vibration due to fluid forces in wear rings",
            "Install swirl breaks at the seal inlet and tighten radial running clearances",
        ],
        [
            "Suction strainer clogged with magnetite scale causing cavitation and pressure drop",
            "Clean and pressure test suction strainer; implement daily delta-P monitoring",
        ],
        [
            "Impeller honeycomb pitting erosion from low-flow cavitation",
            "Upgrade to OEM 13Cr martensitic stainless steel impeller and maintain stable flow",
        ],
    ]

    report_input = XlsxReportInput(
        filename=filename,
        title="Pump Problems and Recommended Improvements",
        headers=headers,
        rows=rows,
    )
    result = await execute_xlsx(report_input)
    assert result["filename"] == filename
    assert result["row_count"] == 3

    # 1. Verify directly with openpyxl
    file_path = tmp_path / filename
    assert file_path.exists()
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    assert ws.cell(row=4, column=1).value == "Problem"
    assert ws.cell(row=4, column=2).value == "Recommended Improvement"
    # Row 5 (data row 1)
    assert "vibration" in str(ws.cell(row=5, column=1).value).lower()
    assert "swirl break" in str(ws.cell(row=5, column=2).value).lower()
    wb.close()

    # 2. Verify with artifact_verifier
    verifier_input = ArtifactVerifierInput(
        relative_path=filename,
        expected_columns=headers,
        min_row_count=2,
        expected_content=["vibration", "strainer", "impeller"],
    )
    ver_res = await verify_artifact(verifier_input)
    assert ver_res["verified"] is True
    assert ver_res["status"] == "PASSED_VERIFICATION"
    assert ver_res["row_count"] == 3
    assert ver_res["column_count"] == 2
