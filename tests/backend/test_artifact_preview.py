from pathlib import Path
import io
import openpyxl
from docx import Document
import pytest
from fastapi.testclient import TestClient

from backend.main import app
from backend.config import settings


@pytest.fixture
def client(tmp_path: Path, monkeypatch):
    """TestClient with sandbox_dir redirected to tmp_path."""
    monkeypatch.setattr(settings, "sandbox_dir", tmp_path)
    return TestClient(app), tmp_path


def test_preview_txt_file(client):
    test_client, sb_dir = client
    txt_file = sb_dir / "notes.txt"
    txt_file.write_text("MRPL Refinery Maintenance Log: All pumps normal.", encoding="utf-8")

    resp = test_client.get("/api/artifacts/notes.txt/preview")
    assert resp.status_code == 200
    data = resp.json()
    assert data["type"] == "txt"
    assert "MRPL Refinery Maintenance Log" in data["content"]
    assert data["truncated"] is False
    assert data["filename"] == "notes.txt"


def test_preview_md_file(client):
    test_client, sb_dir = client
    md_file = sb_dir / "report.md"
    md_file.write_text("# Incident Summary\n\n- Pump: P-204\n- Severity: High", encoding="utf-8")

    resp = test_client.get("/api/artifacts/report.md/preview")
    assert resp.status_code == 200
    data = resp.json()
    assert data["type"] == "md"
    assert "# Incident Summary" in data["content"]


def test_preview_docx_file(client):
    test_client, sb_dir = client
    doc = Document()
    doc.add_paragraph("P-204 Inspection Executive Summary")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Parameter"
    t.cell(0, 1).text = "Value"
    t.cell(1, 0).text = "Discharge"
    t.cell(1, 1).text = "42 bar"
    doc_path = sb_dir / "inspection.docx"
    doc.save(str(doc_path))

    resp = test_client.get("/api/artifacts/inspection.docx/preview")
    assert resp.status_code == 200
    data = resp.json()
    assert data["type"] == "docx"
    assert "paragraphs" in data["content"]
    assert "tables" in data["content"]
    assert "P-204 Inspection Executive Summary" in data["content"]["paragraphs"]
    assert len(data["content"]["tables"]) == 1
    assert data["content"]["tables"][0]["headers"] == ["Parameter", "Value"]


def test_preview_xlsx_file(client):
    test_client, sb_dir = client
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lab_Results"
    ws.append(["Sample_ID", "Stream", "Sulfur_ppm", "Status"])
    ws.append(["SMP-001", "Diesel", "8.2", "PASS"])
    ws.append(["SMP-002", "Naphtha", "4.1", "PASS"])
    xlsx_path = sb_dir / "lab_results.xlsx"
    wb.save(str(xlsx_path))
    wb.close()

    resp = test_client.get("/api/artifacts/lab_results.xlsx/preview")
    assert resp.status_code == 200
    data = resp.json()
    assert data["type"] == "xlsx"
    assert data["content"]["headers"] == ["Sample_ID", "Stream", "Sulfur_ppm", "Status"]
    assert len(data["content"]["rows"]) == 2
    assert data["content"]["rows"][0][0] == "SMP-001"
    assert data["content"]["sheet_name"] == "Lab_Results"


def test_preview_missing_file_404(client):
    test_client, _ = client
    resp = test_client.get("/api/artifacts/nonexistent_file_999.txt/preview")
    assert resp.status_code == 404


def test_preview_path_traversal_rejected(client):
    test_client, _ = client
    resp = test_client.get("/api/artifacts/..%2F..%2Fetc%2Fpasswd/preview")
    assert resp.status_code in (400, 404)


def test_preview_content_size_limit(client):
    test_client, sb_dir = client
    large_file = sb_dir / "huge.txt"
    # Write 600 KB of text
    large_file.write_text("A" * (600 * 1024), encoding="utf-8")

    resp = test_client.get("/api/artifacts/huge.txt/preview")
    assert resp.status_code == 200
    data = resp.json()
    assert data["truncated"] is True
    assert len(data["content"]) <= 500 * 1024 + 10


def test_preview_corrupt_file_422(client):
    test_client, sb_dir = client
    corrupt_docx = sb_dir / "corrupt.docx"
    corrupt_docx.write_bytes(b"PK\x03\x04not a valid zip file at all")

    resp = test_client.get("/api/artifacts/corrupt.docx/preview")
    assert resp.status_code == 422
    assert "Failed to parse" in resp.json().get("detail", "")


def test_preview_unsupported_format(client):
    test_client, sb_dir = client
    bin_file = sb_dir / "firmware.bin"
    bin_file.write_bytes(b"\x00\x01\x02\x03\x04")

    resp = test_client.get("/api/artifacts/firmware.bin/preview")
    assert resp.status_code == 200
    data = resp.json()
    assert data["type"] == "unsupported"
    assert data["content"] is None
