import io
import pytest
from docx import Document
from backend.rag.ingest import DocumentParser, TextChunker


def _create_docx(paragraphs_and_tables):
    """
    Helper to generate DOCX bytes in-memory.
    paragraphs_and_tables is a list of tuples:
      - ("p", "text")
      - ("t", [["h1", "h2"], ["r1c1", "r1c2"]])
    """
    doc = Document()
    for item_type, data in paragraphs_and_tables:
        if item_type == "p":
            doc.add_paragraph(data)
        elif item_type == "t":
            table = doc.add_table(rows=len(data), cols=len(data[0]))
            for r_idx, row in enumerate(data):
                for c_idx, val in enumerate(row):
                    table.cell(r_idx, c_idx).text = str(val)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_with_single_table():
    data = [
        ("p", "Header description"),
        ("t", [["Item", "Value"], ["Pressure", "150 psi"], ["Temp", "75 C"]]),
    ]
    docx_bytes = _create_docx(data)
    text, meta = DocumentParser._parse_docx(docx_bytes)
    assert "Header description" in text
    assert "| Item | Value |" in text
    assert "| Pressure | 150 psi |" in text
    assert "| Temp | 75 C |" in text
    assert meta.get("table_count") == 1
    assert meta.get("has_tables") is True


def test_docx_with_multiple_tables():
    data = [
        ("p", "Section 1"),
        ("t", [["ColA", "ColB"], ["1", "2"]]),
        ("p", "Section 2"),
        ("t", [["ColX", "ColY"], ["3", "4"]]),
    ]
    docx_bytes = _create_docx(data)
    text, meta = DocumentParser._parse_docx(docx_bytes)
    assert "| ColA | ColB |" in text
    assert "| ColX | ColY |" in text
    assert meta.get("table_count") == 2


def test_docx_paragraph_and_table_interleaved():
    data = [
        ("p", "Before Table 1"),
        ("t", [["T1_H", "T1_V"], ["A", "B"]]),
        ("p", "Middle paragraph"),
        ("t", [["T2_H", "T2_V"], ["C", "D"]]),
        ("p", "After Table 2"),
    ]
    docx_bytes = _create_docx(data)
    text, _ = DocumentParser._parse_docx(docx_bytes)

    idx_p1 = text.find("Before Table 1")
    idx_t1 = text.find("T1_H")
    idx_p2 = text.find("Middle paragraph")
    idx_t2 = text.find("T2_H")
    idx_p3 = text.find("After Table 2")

    assert idx_p1 < idx_t1 < idx_p2 < idx_t2 < idx_p3


def test_docx_table_empty_cells():
    data = [
        ("t", [["Col1", "Col2"], ["", "Val2"], ["Val3", ""]]),
    ]
    docx_bytes = _create_docx(data)
    text, meta = DocumentParser._parse_docx(docx_bytes)
    assert "| Col1 | Col2 |" in text
    assert "|  | Val2 |" in text or "| | Val2 |" in text
    assert meta.get("has_tables") is True


def test_docx_table_headers_preserved():
    headers = ["Equipment Tag", "Design Flow", "Discharge Pressure", "Status"]
    data = [
        ("t", [headers, ["P-204", "120 m3/h", "42 bar", "OPERATIONAL"]]),
    ]
    docx_bytes = _create_docx(data)
    text, _ = DocumentParser._parse_docx(docx_bytes)
    for h in headers:
        assert h in text
    assert "---" in text  # markdown table divider


def test_docx_paragraphs_only_unchanged():
    data = [
        ("p", "Paragraph one."),
        ("p", "Paragraph two."),
    ]
    docx_bytes = _create_docx(data)
    text, meta = DocumentParser._parse_docx(docx_bytes)
    assert "Paragraph one." in text
    assert "Paragraph two." in text
    assert meta.get("table_count") == 0
    assert meta.get("has_tables") is False


def test_docx_table_content_retrievable():
    from backend.rag.ingest import Document
    data = [
        ("p", "MRPL Refinery Inspection Runbook for Hydrocracker P-204"),
        ("t", [["Component", "Serial", "Vibration Limit"], ["Impeller", "IMP-9921", "4.5 mm/s"]]),
    ]
    docx_bytes = _create_docx(data)
    text, _ = DocumentParser._parse_docx(docx_bytes)
    doc_obj = Document(document_id="d1", filename="test.docx", file_type="docx", text=text)
    chunker = TextChunker(chunk_size=500, chunk_overlap=50)
    chunks = chunker.chunk(doc_obj)
    assert len(chunks) > 0
    full_chunked_text = " ".join(c.text for c in chunks)
    assert "IMP-9921" in full_chunked_text
    assert "Vibration Limit" in full_chunked_text
